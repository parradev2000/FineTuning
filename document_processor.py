"""
Módulo para procesar diferentes tipos de documentos
"""
import json
import logging
import os

import PyPDF2
from docx import Document

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {'.txt', '.pdf', '.docx', '.json'}


def extract_text_from_pdf(file_path):
    """Extrae texto de un archivo PDF"""
    text = ""
    try:
        with open(file_path, 'rb') as file:
            try:
                pdf_reader = PyPDF2.PdfReader(file)
            except PyPDF2.errors.PdfReadError as e:
                logger.error("PDF corrupto o ilegible %s: %s", file_path, e)
                return text

            num_pages = len(pdf_reader.pages)
            logger.info(
                "Procesando PDF %s con %d página(s)", file_path, num_pages
            )

            for page_num, page in enumerate(pdf_reader.pages, start=1):
                try:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
                    else:
                        logger.warning(
                            "No se pudo extraer texto de la página %d en %s",
                            page_num,
                            file_path,
                        )
                except Exception as e:
                    logger.warning(
                        "Error al extraer texto de la página %d en %s: %s",
                        page_num,
                        file_path,
                        e,
                    )
    except FileNotFoundError:
        logger.error("Archivo PDF no encontrado: %s", file_path)
    except PermissionError:
        logger.error("Permiso denegado al leer PDF: %s", file_path)
    except Exception as e:
        logger.error("Error inesperado al leer PDF %s: %s", file_path, e)
    return text


def extract_text_from_docx(file_path):
    """Extrae texto de un archivo DOCX"""
    text = ""
    try:
        doc = Document(file_path)
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        logger.info(
            "DOCX %s procesado correctamente (%d párrafo(s))",
            file_path,
            len(doc.paragraphs),
        )
    except FileNotFoundError:
        logger.error("Archivo DOCX no encontrado: %s", file_path)
    except PermissionError:
        logger.error("Permiso denegado al leer DOCX: %s", file_path)
    except Exception as e:
        logger.error("Error al leer DOCX %s: %s", file_path, e)
    return text


def _read_text_file(file_path):
    """Lee un archivo de texto plano con manejo de codificación."""
    encodings = ['utf-8', 'latin-1', 'cp1252']
    for encoding in encodings:
        try:
            with open(file_path, 'r', encoding=encoding) as f:
                content = f.read()
            logger.info(
                "Archivo TXT %s leído correctamente (codificación: %s)",
                file_path,
                encoding,
            )
            return content
        except UnicodeDecodeError:
            logger.warning(
                "No se pudo decodificar %s con %s, intentando siguiente "
                "codificación",
                file_path,
                encoding,
            )
        except FileNotFoundError:
            logger.error("Archivo TXT no encontrado: %s", file_path)
            return ""
        except PermissionError:
            logger.error("Permiso denegado al leer TXT: %s", file_path)
            return ""
        except OSError as e:
            logger.error("Error de E/S al leer TXT %s: %s", file_path, e)
            return ""

    logger.error(
        "No se pudo decodificar %s con ninguna codificación soportada",
        file_path,
    )
    return ""


def _read_json_file(file_path):
    """Lee y valida un archivo JSON, devolviendo los datos de entrenamiento."""
    training_data = []
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
    except FileNotFoundError:
        logger.error("Archivo JSON no encontrado: %s", file_path)
        return training_data
    except PermissionError:
        logger.error("Permiso denegado al leer JSON: %s", file_path)
        return training_data
    except json.JSONDecodeError as e:
        logger.error("JSON inválido en %s: %s", file_path, e)
        return training_data
    except OSError as e:
        logger.error("Error de E/S al leer JSON %s: %s", file_path, e)
        return training_data

    if isinstance(data, list):
        for idx, item in enumerate(data):
            if isinstance(item, str):
                if item.strip():
                    training_data.append({'text': item})
                else:
                    logger.warning(
                        "Elemento vacío en índice %d de %s, omitiendo",
                        idx,
                        file_path,
                    )
            elif isinstance(item, dict) and 'text' in item:
                if isinstance(item['text'], str) and item['text'].strip():
                    training_data.append(item)
                else:
                    logger.warning(
                        "Campo 'text' vacío o no es cadena en índice %d de "
                        "%s, omitiendo",
                        idx,
                        file_path,
                    )
            else:
                logger.warning(
                    "Formato inesperado en índice %d de %s: se esperaba "
                    "cadena o dict con 'text'",
                    idx,
                    file_path,
                )
    elif isinstance(data, dict) and 'text' in data:
        if isinstance(data['text'], str) and data['text'].strip():
            training_data.append(data)
        else:
            logger.warning(
                "Campo 'text' vacío o no es cadena en %s", file_path
            )
    else:
        logger.warning(
            "Estructura JSON no reconocida en %s: se esperaba lista o dict "
            "con 'text'",
            file_path,
        )

    logger.info(
        "JSON %s procesado: %d entrada(s) de entrenamiento extraída(s)",
        file_path,
        len(training_data),
    )
    return training_data


def process_documents(file_paths):
    """
    Procesa los documentos subidos y los convierte en un formato adecuado
    para fine-tuning.
    """
    training_data = []

    if not file_paths:
        logger.warning("No se proporcionaron archivos para procesar")
        return training_data

    for file_path in file_paths:
        try:
            if not os.path.isfile(file_path):
                logger.error(
                    "El archivo no existe o no es un archivo válido: %s",
                    file_path,
                )
                continue

            ext = os.path.splitext(file_path)[1].lower()

            if ext not in SUPPORTED_EXTENSIONS:
                logger.warning(
                    "Extensión no soportada '%s' para el archivo %s, "
                    "omitiendo",
                    ext,
                    file_path,
                )
                continue

            if ext == '.txt':
                content = _read_text_file(file_path)
                if not content or not content.strip():
                    logger.warning(
                        "Archivo TXT vacío o sin contenido útil: %s",
                        file_path,
                    )
                    continue
                segments = split_text(content)
                for segment in segments:
                    training_data.append({'text': segment})

            elif ext == '.pdf':
                content = extract_text_from_pdf(file_path)
                if not content or not content.strip():
                    logger.warning(
                        "No se extrajo texto del PDF: %s", file_path
                    )
                    continue
                segments = split_text(content)
                for segment in segments:
                    training_data.append({'text': segment})

            elif ext == '.docx':
                content = extract_text_from_docx(file_path)
                if not content or not content.strip():
                    logger.warning(
                        "No se extrajo texto del DOCX: %s", file_path
                    )
                    continue
                segments = split_text(content)
                for segment in segments:
                    training_data.append({'text': segment})

            elif ext == '.json':
                json_data = _read_json_file(file_path)
                training_data.extend(json_data)

        except Exception as e:
            logger.error(
                "Error inesperado al procesar %s: %s", file_path, e,
                exc_info=True,
            )
            continue

    logger.info(
        "Procesamiento completo: %d archivo(s) procesado(s), %d entrada(s) "
        "de entrenamiento generada(s)",
        len(file_paths),
        len(training_data),
    )
    return training_data


def split_text(text, max_length=1024):
    """
    Divide un texto largo en segmentos manejables
    """
    if not text:
        return []

    if not isinstance(max_length, int) or max_length <= 0:
        logger.warning(
            "max_length inválido (%s), usando valor por defecto 1024",
            max_length,
        )
        max_length = 1024

    segments = []
    for i in range(0, len(text), max_length):
        segment = text[i:i+max_length]
        if segment.strip():  # Solo añadir segmentos que no estén vacíos
            segments.append(segment)
    return segments
